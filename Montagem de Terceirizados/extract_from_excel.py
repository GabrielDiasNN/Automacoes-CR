# {
#   "version": "1.1.0",
#   "skill": "ai-native-development-standard",
#   "contract": "hybrid-fetch-logic",
#   "description": "Extrai dados do Excel com tipagem estrita e logs Base64",
#   "reliability": "Base64-Bridge-Logs"
# }
import base64
import json
import os
import sys
from datetime import datetime
from typing import Any, Dict, List, Optional, cast

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

# Forca UTF-8 para garantir interoperabilidade
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


def log(message: str, level: str = "INFO", exec_id: str = "manual") -> None:
    """Envia logs em Base64 para o stderr (Isolamento total)."""
    ts: str = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    raw_msg: str = f"[{ts}] [PY-EXTRACT-XL] [{level}] [ExecId:{exec_id}] {message}"
    b64_msg: str = base64.b64encode(raw_msg.encode("utf-8")).decode("ascii")
    sys.stderr.write(f"B64:{b64_msg}\n")
    sys.stderr.flush()


def extract() -> None:
    """Funcao principal de extracao de dados do Excel."""
    # pylint: disable=too-many-locals
    exec_id: str = sys.argv[1] if len(sys.argv) > 1 else "manual"
    script_dir: str = os.path.dirname(os.path.abspath(__file__))
    excel_path: str = os.path.join(
        script_dir, "Legacy", "Validador_Notas_Montagem.xlsm"
    )
    table_name: str = "VW_EXC_OB_PED_ROM_Faccao"

    if not os.path.exists(excel_path):
        log(f"Arquivo Excel nao encontrado: {excel_path}", "ERROR", exec_id)
        sys.exit(1)

    try:
        log(f"Abrindo workbook para leitura: {excel_path}", "INFO", exec_id)
        wb = load_workbook(excel_path, data_only=True)

        target_sheet: Optional[Worksheet] = None
        ref: Optional[str] = None

        for sheet in wb.worksheets:
            if table_name in sheet.tables:
                target_sheet = cast(Worksheet, sheet)
                ref = sheet.tables[table_name].ref
                break

        if not target_sheet or not ref:
            log(f"Tabela '{table_name}' nao encontrada.", "ERROR", exec_id)
            sys.exit(1)

        rows = list(target_sheet[ref])

        if not rows:
            log("Tabela vazia ou sem cabecalho.", "ERROR", exec_id)
            sys.exit(1)

        header_cells = rows[0]
        headers: List[str] = [str(cell.value).strip().upper() for cell in header_cells]

        data: List[Dict[str, Any]] = []
        for row in rows[1:]:
            record: Dict[str, Any] = {}
            for i, cell in enumerate(row):
                val: Any = cell.value
                col_name: str = headers[i]
                if isinstance(val, datetime):
                    val = val.isoformat()
                elif isinstance(val, str):
                    val = val.strip()
                record[col_name] = val
            data.append(record)

        # Saida estruturada via STDOUT - EXCLUSIVA PARA DADOS
        sys.stdout.write(json.dumps(data, ensure_ascii=False))
        sys.stdout.flush()
        log(f"Extracao hibrida concluida: {len(data)} registros.", "INFO", exec_id)

    except FileNotFoundError as e:
        log(f"Erro de arquivo: {e}", "ERROR", exec_id)
        sys.exit(1)
    except Exception as e:  # pylint: disable=broad-exception-caught
        log(f"Erro fatal: {e}", "ERROR", exec_id)
        sys.exit(1)


if __name__ == "__main__":
    extract()


